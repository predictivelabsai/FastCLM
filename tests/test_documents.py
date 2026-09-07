import io
import hashlib
from dataclasses import replace

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from fastclm.config import settings
from fastclm.services import documents
from fastclm.services.contracts import ContractService
from fastclm.services.documents import DocumentService, extract_text, file_checksum, safe_filename
from fastclm.services.identity import IdentityService


def test_text_document_ingestion_creates_blocks_attachment_and_version(workspace, tmp_path, monkeypatch):
    actor, _, _ = workspace
    service = ContractService()
    contract = service.create(actor, {"title": "Uploaded agreement"})
    monkeypatch.setattr(documents, "settings", replace(settings, upload_dir=tmp_path / "uploads"))
    result = DocumentService().ingest(actor, contract["id"], "../Agreement.txt", b"1. SERVICES:\n\nThe supplier provides support.\n\nTermination is on 30 days notice.")
    item = service.get(actor, contract["id"])
    assert result["source_filename"] == "Agreement.txt"
    assert result["source_checksum"] == hashlib.sha256(
        b"1. SERVICES:\n\nThe supplier provides support.\n\nTermination is on 30 days notice."
    ).hexdigest()
    assert result["version_number"] == 2
    assert len(item["blocks"]) == 3
    stored = tmp_path / "uploads" / result["storage_path"]
    assert stored.is_file()
    assert file_checksum(stored) == result["source_checksum"]
    stored.write_bytes(b"tampered")
    assert file_checksum(stored) != result["source_checksum"]


def test_document_rejects_unsupported_or_empty_files():
    with pytest.raises(ValueError):
        extract_text("agreement.exe", b"not a contract")
    with pytest.raises(ValueError):
        extract_text("agreement.txt", b"")
    assert safe_filename("../../client/nda?.pdf") == "nda_.pdf"


def test_word_and_text_layer_pdf_extraction():
    word = Document()
    word.add_heading("Services", level=1)
    word.add_paragraph("The supplier provides managed support.")
    word_bytes = io.BytesIO()
    word.save(word_bytes)
    assert "managed support" in extract_text("agreement.docx", word_bytes.getvalue())

    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({
        NameObject("/Type"): NameObject("/Font"),
        NameObject("/Subtype"): NameObject("/Type1"),
        NameObject("/BaseFont"): NameObject("/Helvetica"),
    })
    page[NameObject("/Resources")] = DictionaryObject({
        NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})
    })
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Termination requires thirty days notice.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    pdf_bytes = io.BytesIO()
    writer.write(pdf_bytes)
    assert "Termination requires thirty days notice" in extract_text("agreement.pdf", pdf_bytes.getvalue())


def test_inline_pdf_repeats_session_tenant_and_integrity_checks(fresh_db, tmp_path, monkeypatch):
    from starlette.testclient import TestClient
    import web_app

    upload_dir = tmp_path / "uploads"
    configured = replace(settings, upload_dir=upload_dir)
    monkeypatch.setattr(documents, "settings", configured)
    monkeypatch.setattr(web_app, "settings", configured)
    with TestClient(web_app.app) as client:
        client.post("/signup", data={"name": "Taylor", "organisation": "Taylor Studio", "email": "taylor-inline@example.test", "password": "Secure-password1!"})
        user = fresh_db.one("SELECT id FROM users WHERE email='taylor-inline@example.test'")
        organisation = IdentityService().memberships(user["id"])[0]
        actor = IdentityService().actor(user["id"], organisation["organisation_id"])
        contract = ContractService().create(actor, {"title": "PDF agreement"})
        version = DocumentService().ingest(actor, contract["id"], "agreement.pdf", _pdf_with_text())
        response = client.get(f"/versions/{version['id']}/inline")
        assert response.status_code == 200
        assert response.headers["content-type"] == "application/pdf"
        assert response.headers["content-disposition"].startswith("inline;")
        provenance = client.post("/pdf-provenance", json={
            "source": f"/versions/{version['id']}/inline",
            "evidence": "agreement has searchable text",
        })
        assert provenance.status_code == 200
        assert provenance.json() == {
            "ok": True, "verified": True, "start_word": 1,
            "end_word": 4, "word_count": 4,
        }
        assert client.post("/pdf-provenance", json={
            "source": f"/versions/{version['id']}/inline", "evidence": "invented clause",
        }).json()["verified"] is False

        other_user, other_org = IdentityService().create_workspace("other-inline@example.test", "Secure-password2!", "Other", "Other Studio")
        other_actor = IdentityService().actor(other_user["id"], other_org["id"])
        other_contract = ContractService().create(other_actor, {"title": "Other PDF"})
        other_version = DocumentService().ingest(other_actor, other_contract["id"], "other.pdf", _pdf_with_text())
        assert client.get(f"/versions/{other_version['id']}/inline").status_code == 404
        assert client.post("/pdf-provenance", json={
            "source": f"/versions/{other_version['id']}/inline", "evidence": "searchable text",
        }).status_code == 404

        (upload_dir / version["storage_path"]).write_bytes(b"tampered")
        assert client.get(f"/versions/{version['id']}/inline").status_code == 409


def _pdf_with_text() -> bytes:
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (This agreement has searchable text.) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    content = io.BytesIO()
    writer.write(content)
    return content.getvalue()
